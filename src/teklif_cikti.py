"""Musteriye gidecek teklif ciktisi: PDF (reportlab) ve Excel (openpyxl).

Ic bilgiler (maliyet, kar marji, rakip karsilastirmasi, adam/saat ucreti) ciktiya YAZILMAZ.
"""
import os
from datetime import datetime, timedelta

_FONT_ADAYLARI = [
    ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf'),
    ('C:/Windows/Fonts/segoeui.ttf', 'C:/Windows/Fonts/segoeuib.ttf'),
    ('C:/Windows/Fonts/calibri.ttf', 'C:/Windows/Fonts/calibrib.ttf'),
    ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
    ('/Library/Fonts/Arial Unicode.ttf', '/Library/Fonts/Arial Unicode.ttf'),
]
_ASCII = str.maketrans('çğıöşüÇĞİÖŞÜ', 'cgiosuCGIOSU')


def para(x):
    """1234567.8 -> '1.234.567,80 TL'"""
    s = '{:,.2f}'.format(float(x or 0)).replace(',', 'X').replace('.', ',').replace('X', '.')
    return s + ' TL'


def _tarih(girdi):
    t = girdi.get('tarih')
    try:
        return datetime.strptime(t, '%Y-%m-%d') if t else datetime.now()
    except ValueError:
        return datetime.now()


def _fontlar():
    """Turkce harfleri basabilen TTF kaydeder. Bulamazsa Helvetica + Turkce harfsiz metin."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for normal, kalin in _FONT_ADAYLARI:
        if os.path.exists(normal) and os.path.exists(kalin):
            try:
                pdfmetrics.registerFont(TTFont('TeklifFont', normal))
                pdfmetrics.registerFont(TTFont('TeklifFont-Kalin', kalin))
                return 'TeklifFont', 'TeklifFont-Kalin', None
            except Exception:
                continue
    return 'Helvetica', 'Helvetica-Bold', _ASCII


def satir_aciklamasi(s):
    if s['tur'].lower().startswith(('mas', 'mob')):
        return '%d kullanıcı' % s['adet']
    return '%d adet' % s['adet'] if s['adet'] != 1 else ''


def pdf_olustur(yol, girdi, sonuc, ayar):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether

    font, kalin, cevir = _fontlar()

    def t(metin):
        metin = str(metin)
        return metin.translate(cevir) if cevir else metin

    def P(metin, boyut=9, kalin_mi=False, hiza=None, renk=colors.black):
        stil = ParagraphStyle('p', fontName=kalin if kalin_mi else font, fontSize=boyut, leading=boyut * 1.3,
                              textColor=renk, alignment=hiza if hiza is not None else 0)
        return Paragraph(t(metin).replace('&', '&amp;').replace('<', '&lt;').replace('\n', '<br/>'), stil)

    lacivert = colors.HexColor('#2c3e50')
    acik = colors.HexColor('#ecf0f1')
    kdv = sonuc['kdv_orani']
    tarih = _tarih(girdi)
    gecerlilik = tarih + timedelta(days=int(ayar.get('gecerlilik_gun') or 0))
    belge = SimpleDocTemplate(yol, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                              bottomMargin=16 * mm, title=t('Teklif ' + str(girdi.get('teklif_no', ''))),
                              author=t(ayar.get('firma_adi', '')))
    akis = []

    # ---- Baslik: logo + firma
    firma = [P(ayar.get('firma_adi', ''), 15, True, renk=lacivert)]
    for satir in (ayar.get('firma_adres'), ayar.get('firma_iletisim')):
        if satir:
            firma.append(P(satir, 8, renk=colors.HexColor('#555555')))
    logo = ayar.get('logo_yolu')
    if logo and os.path.exists(logo):
        try:
            resim = Image(logo)
            oran = min(45 * mm / resim.imageWidth, 20 * mm / resim.imageHeight)
            resim.drawWidth, resim.drawHeight = resim.imageWidth * oran, resim.imageHeight * oran
            baslik = Table([[resim, firma]], colWidths=[50 * mm, None])
        except Exception:
            baslik = Table([[firma]])
    else:
        baslik = Table([[firma]])
    baslik.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
    akis += [baslik, Spacer(1, 4 * mm)]
    akis.append(Table([[P('DİA ERP LİSANS + HİZMET TEKLİFİ', 13, True, renk=colors.white)]], colWidths=[None],
                      style=[('BACKGROUND', (0, 0), (-1, -1), lacivert), ('TOPPADDING', (0, 0), (-1, -1), 5),
                             ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    akis.append(Spacer(1, 3 * mm))

    bilgi = [[P('Teklif No', 9, True), P(girdi.get('teklif_no', '-')), P('Tarih', 9, True), P(tarih.strftime('%d.%m.%Y'))],
             [P('Müşteri', 9, True), P(girdi.get('musteri', '')), P('Geçerlilik', 9, True),
              P(gecerlilik.strftime('%d.%m.%Y'))],
             [P('Kullanıcı', 9, True), P('%d masaüstü / %d mobil' % (girdi.get('masaustu', 0), girdi.get('mobil', 0))),
              P('Paket', 9, True), P(sonuc['paket_tipi'])]]
    tb = Table(bilgi, colWidths=[24 * mm, 70 * mm, 24 * mm, None])
    tb.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#bbbbbb')),
                            ('BACKGROUND', (0, 0), (0, -1), acik), ('BACKGROUND', (2, 0), (2, -1), acik),
                            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
    akis += [tb, Spacer(1, 5 * mm)]

    # ---- Lisans kalemleri
    sag = TA_RIGHT
    veri = [[P('#', 8, True, renk=colors.white), P('Paket / Modül', 8, True, renk=colors.white),
             P('Kapsam', 8, True, renk=colors.white), P('Liste Tutarı', 8, True, sag, colors.white),
             P('İndirim', 8, True, sag, colors.white), P('Net Tutar', 8, True, sag, colors.white)]]
    for n, s in enumerate(sonuc['satirlar'], 1):
        veri.append([P(n, 8), P('%s\n%s' % (s['ad'], s['kod']), 8), P(satir_aciklamasi(s), 8),
                     P(para(s['toplam']), 8, hiza=sag), P(para(s['indirim']) if s['indirim'] else '-', 8, hiza=sag),
                     P(para(s['net']), 8, hiza=sag)])
    veri.append(['', P('LİSANS TOPLAMI', 9, True), '', P(para(sonuc['lisans']['toplam']), 8, True, sag),
                 P(para(sonuc['lisans']['indirim']) if sonuc['lisans']['indirim'] else '-', 8, True, sag),
                 P(para(sonuc['lisans']['net']), 8, True, sag)])
    tl = Table(veri, colWidths=[7 * mm, 60 * mm, 21 * mm, 31 * mm, 28 * mm, 31 * mm], repeatRows=1)
    tl.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), lacivert), ('BACKGROUND', (0, -1), (-1, -1), acik),
                            ('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc')),
                            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f8f9fa')]),
                            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
    akis += [P('1. Lisans', 11, True, renk=lacivert), Spacer(1, 2 * mm), tl, Spacer(1, 5 * mm)]

    # ---- Ozet
    ozet = [['Lisans', sonuc['lisans']['net']],
            ['Uyarlama / Kurulum Hizmeti', sonuc['uyarlama']['net']]]
    for k in sonuc['sabit']['kalemler']:
        ozet.append([k['ad'], k['net']])
    ozet.append(['ARA TOPLAM', sonuc['ara_toplam']])
    if sonuc['kampanya']['tutar']:
        ozet.append(['Kampanya İndirimi (%s, %%%g)' % (sonuc['kampanya']['ad'], sonuc['kampanya']['yuzde']),
                     -sonuc['kampanya']['tutar']])
    ozet += [['NET TOPLAM (KDV Hariç)', sonuc['net_toplam']], ['KDV (%%%g)' % kdv, sonuc['kdv']],
             ['GENEL TOPLAM (KDV Dahil)', sonuc['genel_toplam']]]
    kalin_satirlar = {'ARA TOPLAM', 'NET TOPLAM (KDV Hariç)', 'GENEL TOPLAM (KDV Dahil)'}
    to = Table([[P(a, 9, a in kalin_satirlar), P(para(b), 9, a in kalin_satirlar, sag)] for a, b in ozet],
               colWidths=[None, 45 * mm])
    stil = [('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc')),
            ('BACKGROUND', (0, -1), (-1, -1), lacivert)]
    to.setStyle(TableStyle(stil))
    to._cellvalues[-1] = [P(ozet[-1][0], 10, True, renk=colors.white), P(para(ozet[-1][1]), 10, True, sag, colors.white)]
    akis.append(KeepTogether([P('2. Teklif Özeti', 11, True, renk=lacivert), Spacer(1, 2 * mm), to]))
    akis.append(Spacer(1, 5 * mm))

    # ---- Odeme plani
    od = [[P('Ödeme Seçeneği', 8, True, renk=colors.white), P('Taksit Tutarı (KDV Dahil)', 8, True, sag, colors.white),
           P('Toplam (KDV Dahil)', 8, True, sag, colors.white)]]
    for o in sonuc['odeme']:
        od.append([P(o['plan'], 9), P(para(o['dahil']), 9, hiza=sag), P(para(o['dahil'] * o['adet']), 9, hiza=sag)])
    tod = Table(od, colWidths=[None, 50 * mm, 50 * mm])
    tod.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), lacivert),
                             ('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc'))]))
    akis.append(KeepTogether([P('3. Ödeme Planı', 11, True, renk=lacivert), Spacer(1, 2 * mm), tod]))
    akis.append(Spacer(1, 5 * mm))

    notlar = '\n'.join(x for x in [ayar.get('teklif_notlari', ''), girdi.get('notlar', '')] if x)
    if notlar:
        akis.append(KeepTogether([P('Notlar ve Şartlar', 10, True, renk=lacivert), Spacer(1, 1 * mm),
                                  P(notlar, 8, renk=colors.HexColor('#444444'))]))
        akis.append(Spacer(1, 8 * mm))

    onay = Table([[P('MÜŞTERİ ONAYI', 10, True, renk=lacivert), ''],
                  [P('Firma Kaşe / İmza:'), P('_______________________________')],
                  [P('Yetkili Ad Soyad:'), P('_______________________________')],
                  [P('Tarih:'), P('_______________________________')]], colWidths=[45 * mm, None])
    onay.setStyle(TableStyle([('TOPPADDING', (0, 1), (-1, -1), 8)]))
    akis.append(KeepTogether([onay]))

    def sayfa_alti(tuval, doc):
        tuval.saveState()
        tuval.setFont(font, 7)
        tuval.setFillColor(colors.HexColor('#888888'))
        tuval.drawString(16 * mm, 9 * mm, t('%s - Teklif %s' % (ayar.get('firma_adi', ''), girdi.get('teklif_no', ''))))
        tuval.drawRightString(A4[0] - 16 * mm, 9 * mm, t('Sayfa %d' % doc.page))
        tuval.restoreState()

    belge.build(akis, onFirstPage=sayfa_alti, onLaterPages=sayfa_alti)
    return yol


def excel_olustur(yol, girdi, sonuc, ayar):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = 'Teklif'
    lacivert = PatternFill('solid', start_color='2C3E50')
    acik = PatternFill('solid', start_color='ECF0F1')
    beyaz = Font(color='FFFFFF', bold=True)
    ince = Border(bottom=Side(style='thin', color='CCCCCC'))
    para_bicimi = '#,##0.00 "TL"'
    tarih = _tarih(girdi)
    for kol, gen in zip('ABCDEF', (5, 46, 16, 18, 16, 18)):
        ws.column_dimensions[kol].width = gen

    ws['A1'] = ayar.get('firma_adi', '')
    ws['A1'].font = Font(size=16, bold=True, color='2C3E50')
    satir = 2
    for x in (ayar.get('firma_adres'), ayar.get('firma_iletisim')):
        if x:
            ws.cell(satir, 1, x).font = Font(size=9, color='555555')
            satir += 1
    satir += 1
    ws.cell(satir, 1, 'DİA ERP LİSANS + HİZMET TEKLİFİ').font = beyaz
    for c in range(1, 7):
        ws.cell(satir, c).fill = lacivert
    satir += 2
    bilgiler = [('Teklif No', girdi.get('teklif_no', '-'), 'Tarih', tarih.strftime('%d.%m.%Y')),
                ('Müşteri', girdi.get('musteri', ''), 'Geçerlilik',
                 (tarih + timedelta(days=int(ayar.get('gecerlilik_gun') or 0))).strftime('%d.%m.%Y')),
                ('Kullanıcı', '%d masaüstü / %d mobil' % (girdi.get('masaustu', 0), girdi.get('mobil', 0)),
                 'Paket', sonuc['paket_tipi'])]
    for a, b, c, d in bilgiler:
        ws.cell(satir, 1, a).font = Font(bold=True)
        ws.cell(satir, 2, b)
        ws.cell(satir, 4, c).font = Font(bold=True)
        ws.cell(satir, 5, d)
        satir += 1
    satir += 1

    ws.cell(satir, 1, '1. Lisans').font = Font(size=12, bold=True, color='2C3E50')
    satir += 1
    for c, b in enumerate(['#', 'Paket / Modül', 'Kapsam', 'Liste Tutarı', 'İndirim', 'Net Tutar'], 1):
        h = ws.cell(satir, c, b)
        h.font, h.fill = beyaz, lacivert
    satir += 1
    for n, s in enumerate(sonuc['satirlar'], 1):
        degerler = [n, '%s (%s)' % (s['ad'], s['kod']), satir_aciklamasi(s), s['toplam'], s['indirim'], s['net']]
        for c, v in enumerate(degerler, 1):
            h = ws.cell(satir, c, v)
            h.border = ince
            if c >= 4:
                h.number_format = para_bicimi
        ws.cell(satir, 2).alignment = Alignment(wrap_text=True)
        satir += 1
    for c, v in enumerate(['', 'LİSANS TOPLAMI', '', sonuc['lisans']['toplam'], sonuc['lisans']['indirim'],
                           sonuc['lisans']['net']], 1):
        h = ws.cell(satir, c, v)
        h.font, h.fill = Font(bold=True), acik
        if c >= 4:
            h.number_format = para_bicimi
    satir += 2

    ws.cell(satir, 1, '2. Teklif Özeti').font = Font(size=12, bold=True, color='2C3E50')
    satir += 1
    ozet = [('Lisans', sonuc['lisans']['net']), ('Uyarlama / Kurulum Hizmeti', sonuc['uyarlama']['net'])]
    ozet += [(k['ad'], k['net']) for k in sonuc['sabit']['kalemler']]
    ozet.append(('ARA TOPLAM', sonuc['ara_toplam']))
    if sonuc['kampanya']['tutar']:
        ozet.append(('Kampanya İndirimi (%s, %%%g)' % (sonuc['kampanya']['ad'], sonuc['kampanya']['yuzde']),
                     -sonuc['kampanya']['tutar']))
    ozet += [('NET TOPLAM (KDV Hariç)', sonuc['net_toplam']), ('KDV (%%%g)' % sonuc['kdv_orani'], sonuc['kdv']),
             ('GENEL TOPLAM (KDV Dahil)', sonuc['genel_toplam'])]
    for a, b in ozet:
        ws.cell(satir, 2, a)
        h = ws.cell(satir, 6, b)
        h.number_format = para_bicimi
        if a.isupper() or a.startswith(('NET', 'GENEL', 'ARA')):
            ws.cell(satir, 2).font = h.font = Font(bold=True)
        satir += 1
    for c in range(1, 7):
        ws.cell(satir - 1, c).fill = lacivert
    ws.cell(satir - 1, 2).font = ws.cell(satir - 1, 6).font = beyaz
    satir += 1

    ws.cell(satir, 1, '3. Ödeme Planı').font = Font(size=12, bold=True, color='2C3E50')
    satir += 1
    for c, b in ((2, 'Ödeme Seçeneği'), (5, 'Taksit (KDV Dahil)'), (6, 'Toplam (KDV Dahil)')):
        h = ws.cell(satir, c, b)
        h.font, h.fill = beyaz, lacivert
    satir += 1
    for o in sonuc['odeme']:
        ws.cell(satir, 2, o['plan'])
        ws.cell(satir, 5, o['dahil']).number_format = para_bicimi
        ws.cell(satir, 6, round(o['dahil'] * o['adet'], 2)).number_format = para_bicimi
        satir += 1
    satir += 1
    notlar = '\n'.join(x for x in [ayar.get('teklif_notlari', ''), girdi.get('notlar', '')] if x)
    if notlar:
        ws.cell(satir, 1, 'Notlar ve Şartlar').font = Font(bold=True, color='2C3E50')
        satir += 1
        for n in notlar.split('\n'):
            ws.cell(satir, 1, n).font = Font(size=9, color='444444')
            satir += 1
        satir += 1
    ws.cell(satir, 1, 'MÜŞTERİ ONAYI').font = Font(bold=True, color='2C3E50')
    for x in ('Firma Kaşe / İmza:', 'Yetkili Ad Soyad:', 'Tarih:'):
        satir += 2
        ws.cell(satir, 1, x)
        ws.cell(satir, 3, '______________________________')
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.save(yol)
    return yol


def ybh_excel(yol, musteri, tip, sonuc, ayar):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Yıllık Bakım'
    ws.append([ayar.get('firma_adi', '')])
    ws.append(['YILLIK BAKIM (YBH) HESABI', '', 'Müşteri:', musteri, 'Paket:', tip])
    ws.append([])
    ws.append(['Paket Kodu', 'Paket Adı', 'Kullanıcı', 'Modül', 'Lisans Bedeli', 'YBH %', 'Baz Bakım'])
    for s in sonuc['satirlar']:
        ws.append([s['kod'], s['ad'], s['kullanici'], s['modul'], s['lisans'], s['ybh_yuzde'], s['baz']])
    ws.append([])
    for a, b in [('Toplam Baz Bakım', sonuc['baz']), ('Modül Katsayısı (%d modül)' % sonuc['modul'],
                                                       sonuc['modul_katsayi']),
                 ('Kullanıcı Katsayısı (%d kullanıcı)' % sonuc['kullanici'], sonuc['kullanici_katsayi']),
                 ('Toplam Katsayı', sonuc['katsayi']), ('YILLIK BAKIM (KDV Hariç)', sonuc['yillik']['net']),
                 ('KDV', sonuc['yillik']['kdv']), ('KDV DAHİL TOPLAM', sonuc['yillik']['dahil'])]:
        ws.append(['', a, '', '', '', '', b])
    for r in ws.iter_rows(min_row=1, max_row=4):
        for h in r:
            h.font = Font(bold=True)
    for kol, gen in zip('ABCDEFG', (22, 40, 10, 8, 16, 8, 16)):
        ws.column_dimensions[kol].width = gen
    wb.save(yol)
    return yol
