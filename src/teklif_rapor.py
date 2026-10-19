"""Satis rapor panosu: kazanma oranlari, asama hunisi, aylik trend, satisci, kayip nedenleri, indirim etkisi,
en cok teklif edilen paketler. Girdi TeklifVeri.rapor_verisi() satirlaridir; arayuzden bagimsizdir (test edilir)."""
from datetime import datetime

from teklif_hesap import katla

INDIRIM_BANTLARI = [(0, 5), (5, 10), (10, 15), (15, 20), (20, 101)]


def _oran(kaz, kay):
    return round(kaz * 100.0 / (kaz + kay)) if kaz + kay else None


def _gun(a, b):
    try:
        return (datetime.strptime(b[:10], '%Y-%m-%d') - datetime.strptime(a[:10], '%Y-%m-%d')).days
    except (TypeError, ValueError):
        return None


def filtrele(satirlar, bas=None, bit=None, hazirlayan=None):
    sonuc = []
    for r in satirlar:
        t = (r.get('tarih') or '')[:10]
        if bas and t < bas:
            continue
        if bit and t > bit:
            continue
        if hazirlayan and (r.get('hazirlayan') or '') != hazirlayan:
            continue
        sonuc.append(r)
    return sonuc


def rapor(satirlar, aktiviteler=None):
    """satirlar: [{no, tarih, musteri, net, sonuc, red_nedeni, asama, hazirlayan, olasilik, sonuc_tarihi, girdi}]"""
    kaz = [r for r in satirlar if r['sonuc'] == 'KAZANDI']
    kay = [r for r in satirlar if r['sonuc'] == 'KAYBETTİ']
    bek = [r for r in satirlar if r['sonuc'] not in ('KAZANDI', 'KAYBETTİ')]
    tutar = lambda liste: round(sum(r.get('net') or 0 for r in liste), 2)  # noqa: E731
    sureler = [d for d in (_gun(r.get('tarih'), r.get('sonuc_tarihi')) for r in kaz) if d is not None]
    kpi = {
        'teklif': len(satirlar), 'teklif_tutar': tutar(satirlar),
        'kazanilan': len(kaz), 'kazanilan_tutar': tutar(kaz),
        'kaybedilen': len(kay), 'kaybedilen_tutar': tutar(kay),
        'bekleyen': len(bek), 'bekleyen_tutar': tutar(bek),
        'kazanma_orani': _oran(len(kaz), len(kay)),
        'kazanma_orani_tutar': _oran(tutar(kaz), tutar(kay)),
        'ortalama_teklif': round(tutar(satirlar) / len(satirlar), 2) if satirlar else 0,
        'ortalama_karar_gunu': round(sum(sureler) / len(sureler)) if sureler else None,
        'agirlikli_pipeline': round(sum((r.get('net') or 0) * (r.get('olasilik') or 0) / 100.0 for r in bek), 2),
    }

    huni = {}
    for r in satirlar:
        a = r.get('asama') or ('Kazandı' if r['sonuc'] == 'KAZANDI' else 'Kaybetti' if r['sonuc'] == 'KAYBETTİ'
                               else 'Teklif verildi')
        h = huni.setdefault(a, {'asama': a, 'adet': 0, 'tutar': 0.0})
        h['adet'] += 1
        h['tutar'] += r.get('net') or 0

    aylik = {}
    for r in satirlar:
        ay = (r.get('tarih') or '')[:7] or '-'
        a = aylik.setdefault(ay, {'ay': ay, 'teklif': 0, 'tutar': 0.0, 'kazanilan': 0, 'kazanilan_tutar': 0.0,
                                  'kaybedilen': 0})
        a['teklif'] += 1
        a['tutar'] += r.get('net') or 0
        if r['sonuc'] == 'KAZANDI':
            a['kazanilan'] += 1
            a['kazanilan_tutar'] += r.get('net') or 0
        elif r['sonuc'] == 'KAYBETTİ':
            a['kaybedilen'] += 1
    for a in aylik.values():
        a['oran'] = _oran(a['kazanilan'], a['kaybedilen'])

    satisci = {}
    for r in satirlar:
        ad = r.get('hazirlayan') or '(belirtilmemiş)'
        x = satisci.setdefault(ad, {'hazirlayan': ad, 'teklif': 0, 'tutar': 0.0, 'kazanilan': 0, 'kaybedilen': 0,
                                    'kazanilan_tutar': 0.0, 'indirim_top': 0.0})
        x['teklif'] += 1
        x['tutar'] += r.get('net') or 0
        x['indirim_top'] += r.get('indirim') or 0
        if r['sonuc'] == 'KAZANDI':
            x['kazanilan'] += 1
            x['kazanilan_tutar'] += r.get('net') or 0
        elif r['sonuc'] == 'KAYBETTİ':
            x['kaybedilen'] += 1
    for x in satisci.values():
        x['oran'] = _oran(x['kazanilan'], x['kaybedilen'])
        x['ort_indirim'] = round(x.pop('indirim_top') / x['teklif'], 1) if x['teklif'] else 0

    nedenler = {}
    for r in kay:
        n = (r.get('red_nedeni') or '(girilmemiş)').strip()
        anahtar = katla(n).split(' ')[0] if n != '(girilmemiş)' else n
        x = nedenler.setdefault(anahtar, {'neden': n, 'adet': 0, 'tutar': 0.0})
        x['adet'] += 1
        x['tutar'] += r.get('net') or 0

    bantlar = []
    for alt, ust in INDIRIM_BANTLARI:
        grup = [r for r in satirlar if alt <= (r.get('indirim') or 0) < ust]
        k = sum(1 for r in grup if r['sonuc'] == 'KAZANDI')
        y = sum(1 for r in grup if r['sonuc'] == 'KAYBETTİ')
        bantlar.append({'bant': '%%%d-%s' % (alt, '%%%d' % ust if ust <= 100 else 'üstü'), 'teklif': len(grup),
                        'kazanilan': k, 'kaybedilen': y, 'oran': _oran(k, y)})

    paketler = {}
    for r in satirlar:
        for k in (r.get('girdi') or {}).get('kalemler', []):
            x = paketler.setdefault(k.get('kod'), {'kod': k.get('kod'), 'teklif': 0, 'kazanilan': 0})
            x['teklif'] += 1
            if r['sonuc'] == 'KAZANDI':
                x['kazanilan'] += 1

    akt = {}
    for a in aktiviteler or []:
        akt[a.get('tur') or '-'] = akt.get(a.get('tur') or '-', 0) + 1

    return {'kpi': kpi,
            'huni': sorted(huni.values(), key=lambda x: _ASAMA_SIRA.get(x['asama'], 99)),
            'aylik': sorted(aylik.values(), key=lambda x: x['ay']),
            'satisci': sorted(satisci.values(), key=lambda x: -x['kazanilan_tutar']),
            'nedenler': sorted(nedenler.values(), key=lambda x: -x['adet']),
            'indirim': bantlar,
            'paketler': sorted(paketler.values(), key=lambda x: -x['teklif'])[:15],
            'aktivite': sorted(akt.items(), key=lambda x: -x[1])}


_ASAMA_SIRA = {a: i for i, a in enumerate(['Fırsat', 'Keşif', 'Teklif verildi', 'Müzakere', 'Sözlü onay',
                                          'Kazandı', 'Kaybetti'])}


def excel_yaz(yol, r, baslik=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Özet'
    kalin, dolgu = Font(bold=True, color='FFFFFF'), PatternFill('solid', start_color='2C3E50')
    ws.append(['SATIŞ RAPORU', baslik])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    k = r['kpi']
    for etiket, deger in [('Teklif sayısı', k['teklif']), ('Teklif tutarı (KDV hr.)', k['teklif_tutar']),
                          ('Kazanılan', k['kazanilan']), ('Kazanılan tutar', k['kazanilan_tutar']),
                          ('Kaybedilen', k['kaybedilen']), ('Kaybedilen tutar', k['kaybedilen_tutar']),
                          ('Bekleyen', k['bekleyen']), ('Bekleyen tutar', k['bekleyen_tutar']),
                          ('Kazanma oranı % (adet)', k['kazanma_orani']),
                          ('Kazanma oranı % (tutar)', k['kazanma_orani_tutar']),
                          ('Ortalama teklif', k['ortalama_teklif']),
                          ('Ortalama karar süresi (gün)', k['ortalama_karar_gunu']),
                          ('Ağırlıklı bekleyen (olasılıkla)', k['agirlikli_pipeline'])]:
        ws.append([etiket, deger])
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 18
    for ad, anahtarlar, basliklar in [
            ('Aşama hunisi', ('asama', 'adet', 'tutar'), ('Aşama', 'Adet', 'Tutar')),
            ('Aylık', ('ay', 'teklif', 'tutar', 'kazanilan', 'kazanilan_tutar', 'kaybedilen', 'oran'),
             ('Ay', 'Teklif', 'Tutar', 'Kazanılan', 'Kazanılan tutar', 'Kaybedilen', 'Kazanma %')),
            ('Satışçı', ('hazirlayan', 'teklif', 'tutar', 'kazanilan', 'kaybedilen', 'kazanilan_tutar', 'oran',
                         'ort_indirim'),
             ('Hazırlayan', 'Teklif', 'Tutar', 'Kazanılan', 'Kaybedilen', 'Kazanılan tutar', 'Kazanma %',
              'Ort. indirim %')),
            ('Kayıp nedenleri', ('neden', 'adet', 'tutar'), ('Neden', 'Adet', 'Tutar')),
            ('İndirim etkisi', ('bant', 'teklif', 'kazanilan', 'kaybedilen', 'oran'),
             ('İndirim', 'Teklif', 'Kazanılan', 'Kaybedilen', 'Kazanma %')),
            ('Paketler', ('kod', 'teklif', 'kazanilan'), ('Paket', 'Teklif', 'Kazanılan'))]:
        s = wb.create_sheet(ad)
        s.append(list(basliklar))
        for h in s[1]:
            h.font, h.fill = kalin, dolgu
        kaynak = {'Aşama hunisi': 'huni', 'Aylık': 'aylik', 'Satışçı': 'satisci', 'Kayıp nedenleri': 'nedenler',
                  'İndirim etkisi': 'indirim', 'Paketler': 'paketler'}[ad]
        for x in r[kaynak]:
            s.append([x.get(a) for a in anahtarlar])
        for i in range(len(basliklar)):
            s.column_dimensions[chr(65 + i)].width = 18 if i else 26
    wb.save(yol)
    return yol
