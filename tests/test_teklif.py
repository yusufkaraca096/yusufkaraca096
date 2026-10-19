"""Teklif motoru testleri (uydurma fiyatlarla; gercek fiyat listesi depoda yoktur)."""
import os
import sqlite3
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

import teklif_cikti  # noqa: E402
from teklif_hesap import hesapla, ybh_hesapla, kademe_katsayi, kod_listesi_coz, sayi_coz  # noqa: E402
from teklif_veri import TeklifVeri, AYAR_VARSAYILAN, varsayilan_klasor, veri_klasoru_degistir  # noqa: E402


def motor_excel(yol):
    """RACA_DIA_Teklif_Motoru.xlsx ile ayni yapida kucuk dosya."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'PaketListesi'
    ws.append(['PAKET LİSTESİ', 'Kodu', 'Adı', 'Std. Fiyat', 'Ent. Fiyat', 'Kullanıcı %', 'YBH %', 'Açıklama',
               'Ürünler', 'Zorunlu'])
    ws.append(['Masaüstü', 'MAS1', 'Muhasebe', 1000, 2000, 10, 15, 'açıklama', 'a-b-c', '[]'])
    ws.append(['Masaüstü', 'MAS2', 'Ticari', 500, 0, 20, 15, '', 'a-b', '[]'])
    ws.append(['Mobil', 'MOB1', 'Mobil Satış', 300, 400, 50, 10, '', 'm', '[]'])
    ws.append(['Diğer', 'DIG1', 'E-Fatura', 800, 900, 0, 20, '', '', '[]'])
    ws.append(['Diğer', 'URE', 'Üretim', 600, 600, 0, 15, '', 'ure', '[a, x]'])
    ws.append(['Masaüstü', 'EKO', 'Ticari Eko', 200, 0, 15, 15, 'Stok-Cari [Max 3 kullanıcı destekler]', 'a', '[]'])
    ws.append(['Diğer', 'XPKT', 'X Modülü', 100, 0, 0, 0, '', 'x-y', '[]'])
    a = wb.create_sheet('ADAM SAAT')
    a.append(['ADAM/SAAT'])
    a.append([])
    a.append([])
    a.append(['Paket Kodu', 'Modül', 'Standart', 'Enterprise'])
    for r, (std, ent) in enumerate([(2, 3), (1, None), (1, 2), (4, 5)], start=5):
        a.cell(r, 1, '=IF(PaketListesi!B%d="","",PaketListesi!B%d)' % (r - 3, r - 3))   # Excel'deki gibi formul
        a.cell(r, 3, std)
        a.cell(r, 4, '=ROUND(C%d*1.5,0)' % r if ent is None else ent)
    s = wb.create_sheet('AYARLAR')
    s.append(['AYARLAR'])
    s.append([])
    s.append(['Ayar', 'Değer', 'Açıklama'])
    for satir in [('Enterprise Masaüstü Eşiği', 10, ''), ('Enterprise Mobil Eşiği', 10, ''),
                  ('Lisans Yüzdesi (%)', 10, ''), ('Saatlik Danışman Ücreti', 100, ''),
                  ('Minimum Uyarlama', 500, 'Alt sınır'), ('Maksimum Uyarlama', 0, ''),
                  ('--- SABİT FİYATLAR ---', None, None), ('B2B Premium Uyarlama', 700, 'Sabit fiyat'),
                  ('KDV Oranı (%)', 20, '')]:
        s.append(list(satir))
    k = wb.create_sheet('KAMPANYALAR')
    k.append(['KAMPANYA'])
    k.append([])
    k.append(['Kampanya Adı', 'İndirim %', 'Durum', 'Açıklama'])
    k.append(['YOK', 0, 'AKTİF', ''])
    k.append(['Rakip Geçiş', 15, 'AKTİF', ''])
    k.append(['Eski', 30, 'PASİF', ''])
    k.append(['Toplu Alım', 20, 'AKTİF', '10+ kullanıcı'])
    r = wb.create_sheet('RAKİP')
    for _ in range(3):
        r.append([])
    r.append(['Firma', 'Adam/Saat', 'Saatlik TL', 'Toplam'])
    r.append(['RACA (Biz)', 8, 100])
    r.append(['Logo', 8, 150])
    r.append(['Mikro', 8, 50])
    g = wb.create_sheet('TEKLİF GEÇMİŞİ')
    for _ in range(3):
        g.append([])
    g.append(['Teklif No', 'Tarih', 'Müşteri', 'Paket', 'Kullanıcı', 'Lisans', 'Uyarlama', 'KDV Dahil', 'Kampanya',
              'Sonuç', 'Red Nedeni'])
    g.append(['ESKI-1', '2025-05-01', 'Eski Müşteri', 'STANDART', '3', 1000, 500, 1800, 'YOK', 'KAZANDI', ''])
    wb.save(yol)


class TeklifTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'veri', 'teklif.db'))
        self.rapor = self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))
        self.ayar = self.v.ayarlar()
        self.pk = self.v.paketler()

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def h(self, **g):
        return hesapla(dict({'masaustu': 3, 'mobil': 2}, **g), self.pk, self.ayar, self.v.sabit_kalemler(),
                       self.v.rakipler())

    # ---------------------------------------------------------------- ice aktarma
    def test_ice_aktarma(self):
        self.assertEqual(self.rapor[0], '7 paket')
        self.assertEqual((self.pk['MAS1']['std_saat'], self.pk['MAS1']['ent_saat']), (2, 3))
        self.assertEqual(self.pk['MAS2']['ent_saat'], 2)        # =ROUND(C*1.5) formulu
        self.assertEqual(self.ayar['saatlik_ucret'], 100)
        self.assertEqual(self.v.sabit_kalemler(), {'B2B Premium Uyarlama': 700})
        self.assertEqual([k['ad'] for k in self.v.kampanyalar()], ['YOK', 'Rakip Geçiş', 'Toplu Alım'])  # PASIF yok
        self.assertEqual({k['ad']: k['kosul'] for k in self.v.kampanyalar()}['Toplu Alım'], 'kullanici>=10')
        self.assertEqual(self.pk['URE']['zorunlu'], 'a, x')
        self.assertEqual([r['firma'] for r in self.v.rakipler()], ['Logo', 'Mikro'])
        self.assertEqual(self.v.teklifler()[0]['no'], 'ESKI-1')

    # ---------------------------------------------------------------- hesap
    def test_kullanici_ve_paket_tipi(self):
        s = self.h(kalemler=[{'kod': 'MAS1'}, {'kod': 'MOB1'}, {'kod': 'DIG1'}])
        self.assertEqual(s['paket_tipi'], 'STANDART')
        mas, mob, dig = s['satirlar']
        self.assertEqual((mas['adet'], mas['toplam']), (3, 1000 + 2 * 1000 * 0.10))
        self.assertEqual((mob['adet'], mob['toplam']), (2, 300 + 300 * 0.5))
        self.assertEqual((dig['adet'], dig['toplam']), (1, 800))
        s = self.h(masaustu=11, kalemler=[{'kod': 'MAS1'}, {'kod': 'MAS2'}])
        self.assertEqual(s['paket_tipi'], 'ENTERPRISE')
        self.assertEqual(s['satirlar'][0]['birim'], 2000)
        self.assertEqual(s['satirlar'][1]['birim'], 500)          # Ent. fiyat 0 -> Std.
        self.assertTrue(any('Enterprise fiyatı yok' in k['mesaj'] for k in s['kontrol']['liste']))

    def test_indirim_tutari_kdv_haric_ve_kdvyi_etkiler(self):
        """Kullanici sikayeti: Ind. Tutar KDV haric toplami ve KDV alt toplamini etkilemiyordu."""
        s = self.h(kalemler=[{'kod': 'DIG1', 'ind_yuzde': 10, 'ind_tutar': 100}])
        self.assertEqual(s['lisans'], {'net': 620.0, 'kdv': 124.0, 'dahil': 744.0, 'toplam': 800.0,
                                       'indirim': 180.0, 'kamp_indirim': 0.0})
        bos = self.h(kalemler=[{'kod': 'DIG1'}])
        self.assertEqual(s['uyarlama']['sinir'], 'ALT SINIR')       # uyarlama iki durumda da 500
        self.assertEqual(round(bos['net_toplam'] - s['net_toplam'], 2), 180)
        self.assertEqual(round(bos['kdv'] - s['kdv'], 2), 36)
        self.assertEqual(round(bos['genel_toplam'] - s['genel_toplam'], 2), 216)

    def test_uyarlama_tum_satirlarin_saati(self):
        """Excel'de SUM(O17:O24) ilk 7 satiri saymiyordu."""
        s = self.h(kalemler=[{'kod': k} for k in ('MAS1', 'MAS2', 'MOB1', 'DIG1')])
        self.assertEqual(s['uyarlama']['saat'], 2 + 1 + 1 + 4)
        self.assertEqual(s['uyarlama']['saat_bedeli'], 800)
        self.assertEqual(s['uyarlama']['net'], round(s['lisans']['net'] * 0.10 + 800, 2))
        s = self.h(kalemler=[{'kod': 'MAS1', 'saat': 0}])     # elle saat
        self.assertEqual(s['uyarlama']['saat'], 0)
        self.assertEqual((s['uyarlama']['net'], s['uyarlama']['sinir']), (500, 'ALT SINIR'))

    def test_kampanya_ve_odeme_plani_net_uzerinden(self):
        """Excel'de odeme plani net yerine kampanya INDIRIM tutarindan hesaplaniyordu."""
        s = self.h(kalemler=[{'kod': 'DIG1'}], sabitler=['B2B Premium Uyarlama'], kampanya='Rakip Geçiş',
                   kampanya_yuzde=15, maliyet=1000)
        ara = s['ara_toplam']
        self.assertEqual(ara, 800 + 500 + 700)                 # lisans + min uyarlama + sabit
        self.assertEqual(s['net_toplam'], ara * 0.85)
        self.assertEqual(s['genel_toplam'], round(ara * 0.85 * 1.2, 2))
        pesin, iki = s['odeme'][0], s['odeme'][1]
        self.assertEqual(pesin['dahil'], round(ara * 0.85 * 0.95 * 1.2, 2))
        self.assertEqual(iki['taksit_tutari'], round(ara * 0.85 / 2 * 1.2, 2))
        self.assertEqual(iki['dahil'], round(ara * 0.85 * 1.2, 2))
        self.assertEqual(s['kar']['brut'], ara * 0.85 - 1000)

    def test_rakip(self):
        s = self.h(kalemler=[{'kod': 'DIG1'}])                  # 4 saat x 100 = 400
        r = s['rakip']
        self.assertEqual(r['biz'], 400)
        self.assertEqual([x['toplam'] for x in r['satirlar']], [600, 200])     # ayni 4 saat x rakip ucreti
        self.assertEqual(r['ortalama'], 400)
        self.ayar['rakip_ayni_saat'] = 0
        self.assertEqual([x['toplam'] for x in self.h(kalemler=[{'kod': 'DIG1'}])['rakip']['satirlar']], [1200, 400])

    def test_bilinmeyen_kod_sirayi_kaydirmaz(self):
        s = self.h(kalemler=[{'kod': 'YOK'}, {'kod': 'DIG1'}])
        self.assertEqual([x['sira'] for x in s['satirlar']], [1])
        self.assertTrue(s['uyarilar'])

    def test_ybh(self):
        self.assertEqual(kademe_katsayi(5, '5:1; 10:1.1; *:1.65'), 1)
        self.assertEqual(kademe_katsayi(7, '5:1; 10:1.1; *:1.65'), 1.1)
        self.assertEqual(kademe_katsayi(70, '5:1; 10:1.1; *:1.65'), 1.65)
        y = ybh_hesapla('STANDART', [{'kod': 'MAS1', 'kullanici': 3}, {'kod': 'MAS2', 'kullanici': 1}],
                        self.pk, self.ayar)
        self.assertEqual(y['modul'], 5)
        self.assertEqual(y['baz'], round(1200 * 0.15 + 500 * 0.15, 2))
        self.assertEqual(y['yillik']['net'], round(y['baz'] * 1 * 1, 2))

    # ---------------------------------------------------------------- kayit / yedek
    def test_teklif_kaydet_ve_numara(self):
        g = {'musteri': 'A', 'masaustu': 1, 'mobil': 0, 'kalemler': [{'kod': 'MAS1'}]}
        s = hesapla(g, self.pk, self.ayar)
        no1 = self.v.teklif_kaydet(g, s)
        no2 = self.v.teklif_kaydet(g, s)
        self.assertTrue(no1.endswith('-0001') and no2.endswith('-0002'))
        self.v.teklif_sonuc_yaz(no1, 'KAZANDI')
        self.v.teklif_kaydet(dict(g, musteri='A2'), s, no1)          # duzenleme sonucu silmez
        t = {x['no']: x for x in self.v.teklifler()}
        self.assertEqual((t[no1]['musteri'], t[no1]['sonuc']), ('A2', 'KAZANDI'))
        self.assertEqual(self.v.teklif_girdisi(no1)['kalemler'], [{'kod': 'MAS1'}])

    def test_kayitlar_kapatip_acinca_kalir_ve_yedeklenir(self):
        """Kullanici sikayeti: kapatip acinca onceki kayitlar siliniyordu."""
        yol = self.v.yol
        g = {'musteri': 'Kalıcı', 'masaustu': 1, 'mobil': 0, 'kalemler': [{'kod': 'MAS1'}]}
        no = self.v.teklif_kaydet(g, hesapla(g, self.pk, self.ayar))
        self.v.kapat()
        self.v = TeklifVeri(yol)
        self.assertIn(no, [t['no'] for t in self.v.teklifler()])
        son = self.v.yedekler()[0][0]
        self.assertIn('kapanis', son)                       # kaydedilen teklif kapanis yedeginde de var
        y = sqlite3.connect(son)
        self.assertEqual(y.execute('SELECT musteri FROM teklif WHERE no=?', (no,)).fetchone()[0], 'Kalıcı')
        y.close()
        # yedekten geri donus: silinen teklif geri gelir
        self.v.yedek_al('test')
        self.v.teklif_sil(no)
        self.v.yedekten_don(self.v.yedekler()[0][0] if 'test' in self.v.yedekler()[0][1] else
                            [y for y, a in self.v.yedekler() if 'test' in a][0])
        self.assertIn(no, [t['no'] for t in self.v.teklifler()])

    def test_yedek_sayisi_sinirli(self):
        import teklif_veri
        with mock.patch.object(teklif_veri, 'YEDEK_SAYISI', 3):
            for n in range(5):
                with mock.patch('teklif_veri.datetime') as dt:
                    from datetime import datetime
                    dt.now.return_value = datetime(2026, 1, 1, 10, 0, n)
                    self.v.yedek_al('n%d' % n)
        self.assertEqual(len(self.v.yedekler()), 3)

    def test_gecersiz_yedek_reddedilir(self):
        bozuk = os.path.join(self.d, 'bozuk.db')
        sqlite3.connect(bozuk).close()
        with self.assertRaises(ValueError):
            self.v.yedekten_don(bozuk)

    def test_veri_exe_yaninda_tutulmaz(self):
        with mock.patch.dict(os.environ, {'APPDATA': os.path.join(self.d, 'appdata')}), \
                mock.patch('os.path.expanduser', return_value=os.path.join(self.d, 'ev')):
            k = varsayilan_klasor()
            self.assertTrue(k.startswith(os.path.join(self.d, 'ev')))
            self.assertNotIn('src', os.path.relpath(k, self.d))
            os.makedirs(k)
            open(os.path.join(k, 'teklif.db'), 'w').write('x')
            hedef = veri_klasoru_degistir(os.path.join(self.d, 'ortak'))
            self.assertTrue(os.path.exists(hedef))
            self.assertEqual(varsayilan_klasor(), os.path.join(self.d, 'ortak'))

    def test_fiyat_listesi_guncelleme_saatleri_korur(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['DIA fiyat listesi'])
        ws.append(['Tür', 'Kodu', 'Adı', 'Std. Fiyat', 'Ent. Fiyat', 'Kullanıcı %', 'YBH %', 'Açıklama', 'Ürünler'])
        ws.append(['Masaüstü', 'MAS1', 'Muhasebe', 1100, 2200, 10, 15, '', 'a-b-c'])
        ws.append(['Diğer', 'YENI1', 'Yeni', 50, 0, 0, 0, '', ''])
        yol = os.path.join(self.d, 'dia.xlsx')
        wb.save(yol)
        r = self.v.paket_listesi_guncelle(yol)
        self.assertEqual((r['fiyati_degisen'], r['yeni']), (['MAS1'], ['YENI1']))
        self.assertEqual(sorted(r['listeden_kalkan']), ['DIG1', 'EKO', 'MAS2', 'MOB1', 'URE', 'XPKT'])
        pk = self.v.paketler()
        self.assertEqual((pk['MAS1']['std_fiyat'], pk['MAS1']['std_saat']), (1100, 2))
        self.assertEqual(pk['MAS2']['aktif'], 0)
        self.assertEqual(len(self.v.paketler(sadece_aktif=True)), 2)

    def test_ayar_varsayilanlari_ticari_deger_icermez(self):
        for ad in ('saatlik_ucret', 'min_uyarlama', 'max_uyarlama'):
            self.assertEqual(AYAR_VARSAYILAN[ad], 0)

    # ---------------------------------------------------------------- cikti
    def test_pdf_ve_excel(self):
        g = {'teklif_no': 'RT-2026-0001', 'musteri': 'Çağlar Şeker İnş. Ltd. Şti.', 'masaustu': 3, 'mobil': 2,
             'kalemler': [{'kod': 'MAS1', 'ind_tutar': 50}, {'kod': 'DIG1'}], 'kampanya': 'Rakip Geçiş',
             'kampanya_yuzde': 15, 'notlar': 'Özel not'}
        s = hesapla(g, self.pk, self.ayar, self.v.sabit_kalemler())
        pdf = teklif_cikti.pdf_olustur(os.path.join(self.d, 't.pdf'), g, s, self.ayar)
        with open(pdf, 'rb') as f:
            self.assertTrue(f.read(5).startswith(b'%PDF'))
        xl = teklif_cikti.excel_olustur(os.path.join(self.d, 't.xlsx'), g, s, self.ayar)
        from openpyxl import load_workbook
        hucreler = [c.value for r in load_workbook(xl).active.iter_rows() for c in r if c.value is not None]
        self.assertIn(s['genel_toplam'], hucreler)
        self.assertIn('Çağlar Şeker İnş. Ltd. Şti.', hucreler)
        ic = ('kâr', 'komisyon', 'rakip ortalama', 'rakiplerden', 'bize kalan', 'maliyet (dia')   # ic bilgiler
        self.assertFalse([h for h in hucreler if isinstance(h, str) and any(x in h.lower() for x in ic)])
        self.assertEqual(teklif_cikti.para(1234567.8), '1.234.567,80 TL')



class KurallarTest(unittest.TestCase):
    """Hizmet kirilimi, kesif, kontroller, odeme, bakim, satis asistani (varsayilan ornek kurallarla)."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))
        self.ayar = self.v.ayarlar()
        self.pk = self.v.paketler()
        self.k = self.v.kurallar()

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def h(self, **g):
        return hesapla(dict({'musteri': 'A', 'masaustu': 3, 'mobil': 2, 'kalemler': [{'kod': 'MAS1'}]}, **g),
                       self.pk, self.ayar, self.v.sabit_kalemler(), self.v.rakipler(), self.k)

    def hz(self, s):
        return {x['kod']: x for x in s['hizmet']['kalemler']}

    def test_hizmet_kirilimi_ve_kesif(self):
        s = self.h(cevaplar={'sube_sayisi': 3, 'uretim': True, 'gecis_programi': 'Netsis', 'cari_sayisi': 3000,
                             'stok_sayisi': 1000, 'rapor_sayisi': 2, 'veri_hata_orani': 10})
        h = self.hz(s)
        self.assertEqual(h['kurulum']['saat'], 2 + 2 * 4 + 24)                  # paket + 2 ek sube + uretim
        self.assertEqual(h['veri_aktarimi']['saat'], round((4000 / 2000 + 4 + 8) * 1.1, 1))
        self.assertEqual(h['egitim']['saat'], 8)                               # 5 kullanici -> 1 grup x 8
        self.assertEqual(h['rapor']['saat'], 8)
        self.assertEqual(h['canli_gecis']['saat'], 8)
        diger = sum(x['saat'] for k, x in h.items() if k != 'proje')
        self.assertEqual(h['proje']['saat'], round(diger * 0.10, 1))
        self.assertEqual(s['uyarlama']['saat'], round(diger + h['proje']['saat'], 1))
        self.assertTrue(any('Netsis' in n for n in h['veri_aktarimi']['nedenler']))

    def test_genel_katsayi_elle_ve_kalibrasyon(self):
        s = self.h(cevaplar={'it_personeli_yok': True})
        self.assertEqual(self.hz(s)['kurulum']['saat'], round(2 * 1.15, 1))
        s = self.h(hizmet_elle={'egitim': 3})
        self.assertEqual(self.hz(s)['egitim']['saat'], 3)
        self.assertEqual(self.hz(s)['egitim']['oneri_saat'], 8)
        self.v.carpan_yaz('kurulum', 1.5)
        self.k = self.v.kurallar()
        self.assertEqual(self.hz(self.h())['kurulum']['saat'], 3)

    def test_saat_ucreti_ve_seyahat(self):
        for ad, d in (('yol_ucreti', 100), ('konaklama_ucreti', 50), ('gundelik', 10), ('ziyaret_basina_gun', 3)):
            self.ayar[ad] = d
        s = self.h(saatlik_ucret=50, seyahat={'sehir_disi': True, 'gun': 5, 'kisi': 2})
        self.assertEqual(s['uyarlama']['saat_ucreti'], 50)
        self.assertEqual(s['seyahat']['tutar'], 2 * (2 * 100 + 3 * 50 + 5 * 10))     # 2 ziyaret, 3 gece
        self.assertIn(s['seyahat']['tutar'], [s['ara_toplam'] - s['lisans']['net'] - s['uyarlama']['net']])

    def test_paket_kontrolleri(self):
        s = self.h(masaustu=5, kalemler=[{'kod': 'EKO'}, {'kod': 'URE'}, {'kod': 'MAS1'}, {'kod': 'MAS2'}])
        m = ' | '.join(k['mesaj'] for k in s['kontrol']['liste'])
        self.assertIn('EKO en fazla 3 kullanıcı', m)
        self.assertIn('URE için gerekli modüller teklifte yok: x', m)
        self.assertIn('XPKT', m)                                               # eksik modulu iceren paket onerisi
        self.assertIn('MAS2 modüllerinin %100', m)                             # cift ucretlendirme
        self.assertTrue(s['kontrol']['kritik'])

    def test_indirim_kar_ve_ucret_sinirlari(self):
        self.ayar.update(max_toplam_indirim=10, min_kar_marji=50, min_saat_ucreti=80)
        s = self.h(kampanya='Rakip Geçiş', kampanya_yuzde=15, maliyet=100000, saatlik_ucret=60)
        m = ' | '.join(k['mesaj'] for k in s['kontrol']['liste'] if k['seviye'] == 'KRİTİK')
        self.assertIn('Toplam indirim', m)
        self.assertIn('Kâr marjı', m)
        self.assertIn('Saatlik ücret', m)

    def test_eksik_bilgi_kontrolleri(self):
        s = self.h(cevaplar={'gecis_programi': 'Logo'})
        m = ' | '.join(k['mesaj'] for k in s['kontrol']['liste'])
        self.assertIn('yetkili', m)
        self.assertIn('kayıt sayısı girilmedi', m)
        self.v.ayar_yaz('fiyat_listesi_tarihi', '2020-01-01')
        self.ayar = self.v.ayarlar()
        self.assertIn('fiyat listesi', ' '.join(k['mesaj'] for k in self.h()['kontrol']['liste']))

    def test_kampanya_kosullari(self):
        from teklif_hesap import kosul_saglanir
        from datetime import date
        b = {'kullanici': 12, 'cevaplar': {'gecis_programi': 'Logo', 'yeni_musteri': True}, 'tarih': date(2026, 12, 20)}
        self.assertTrue(kosul_saglanir('kullanici>=10', b))
        self.assertFalse(kosul_saglanir('kullanici>=20', b))
        self.assertTrue(kosul_saglanir('soru:gecis_programi=Logo|Netsis ve soru:yeni_musteri', b))
        self.assertFalse(kosul_saglanir('soru:gecis_programi=ETA', b))
        self.assertTrue(kosul_saglanir('tarih:12-15..12-31', b))
        self.assertTrue(kosul_saglanir('tarih:12-15..01-10', dict(b, tarih=date(2027, 1, 5))))
        self.assertIsNone(kosul_saglanir('elle', b))
        s = self.h(masaustu=12)
        self.assertTrue(any('Toplu Alım' in o['mesaj'] for o in s['asistan']['oneriler']))
        s = self.h(kampanya='Toplu Alım', kampanya_yuzde=20)
        self.assertTrue(any('koşulu sağlanmıyor' in k['mesaj'] for k in s['kontrol']['liste']))

    def test_odeme_vade_farki_ve_komisyon(self):
        from teklif_hesap import odeme_secenekleri
        o = odeme_secenekleri(1000, 20, [{'ad': 'KK6', 'tur': 'KREDI_KARTI', 'taksit': 6, 'vade_farki': 0,
                                          'komisyon': 5},
                                         {'ad': 'Cek', 'tur': 'TAKSIT', 'taksit': 3, 'aylik_vade': 2}], maliyet=400)
        kk, cek = o
        self.assertEqual(kk['dahil'], 1200)
        self.assertEqual(kk['komisyon_tutari'], 60)
        self.assertEqual(kk['bize_kalan'], 940)
        self.assertEqual(kk['kar'], 540)
        self.assertEqual(kk['komisyon_acigi'], 60)
        # gereken vade farkiyla bize net 1000 kalir
        o2 = odeme_secenekleri(1000, 20, [{'ad': 'KK6', 'tur': 'KREDI_KARTI', 'taksit': 6, 'komisyon': 5,
                                           'vade_farki': kk['gereken_vade_farki']}])[0]
        self.assertAlmostEqual(o2['bize_kalan'], 1000, delta=0.1)
        self.assertEqual(cek['vade_farki'], 4)                                  # %2 x ortalama 2 ay
        self.assertEqual(cek['taksit_tutari'], round(1040 * 1.2 / 3, 2))

    def test_bakim_onerisi(self):
        self.ayar.update(saatlik_ucret=100, destek_esikleri='5; 1000')
        s = self.h(cevaplar={'entegrasyon_sayisi': 1})
        b = s['bakim']
        self.assertEqual(b['onerilen'], 'Standart')
        std = next(p for p in b['paketler'] if p['ad'] == 'Standart')
        self.assertEqual(std['yillik'], round((std['saat'] * 100 + 1 * 800) * 0.95, 2))
        self.assertTrue(any('entegrasyon' in d[0] for d in b['detay']))
        s = self.h(bakim='Standart', cevaplar={'entegrasyon_sayisi': 1})
        self.assertEqual(s['tco']['yillar'][0]['bakim'], std['yillik'])
        self.assertEqual(s['tco']['yillar'][0]['ybh'], 0)
        self.assertEqual(s['tco']['yillar'][1]['ybh'], s['ybh']['yillik']['net'])

    def test_satis_asistani_gecmis(self):
        g = {'musteri': 'A', 'masaustu': 3, 'mobil': 2, 'kalemler': [{'kod': 'MAS1'}]}
        for n, sonuc in enumerate(['KAZANDI', 'KAZANDI', 'KAYBETTİ', 'KAZANDI']):
            no = self.v.teklif_kaydet(dict(g, kampanya_yuzde=10 * n), self.h(kampanya_yuzde=10 * n))
            self.v.teklif_sonuc_yaz(no, sonuc, 'fiyat yüksek' if sonuc == 'KAYBETTİ' else '')
        self.k = self.v.kurallar()
        s = self.h()
        m = next(o['mesaj'] for o in s['asistan']['oneriler'] if o['baslik'] == 'Geçmiş tekliflerden')
        self.assertIn("5 teklifin 4'i", m)                      # + Excel'den gelen ESKI-1 (KAZANDI)
        self.assertIn('Benzer büyüklükteki', m)
        self.assertTrue(5 <= s['asistan']['skor'] <= 95)

    def test_revizyon_ve_gerceklesen(self):
        g = {'musteri': 'A', 'masaustu': 3, 'mobil': 0, 'kalemler': [{'kod': 'MAS1'}]}
        no = self.v.teklif_kaydet(g, self.h())
        self.assertEqual(self.v.revizyon_no(no), no + '-R1')
        self.v.teklif_kaydet(g, self.h(), no + '-R1')
        self.assertEqual(self.v.revizyon_no(no + '-R1'), no + '-R2')
        tahmin = self.v.teklif_girdisi(no)['_ozet']['hizmet']
        self.assertIn('kurulum', tahmin)
        self.v.gerceklesen_yaz(no, {'kurulum': (tahmin['kurulum'], tahmin['kurulum'] * 2)})
        r = self.v.tahmin_gercek_raporu()[0]
        self.assertEqual((r['kalem'], r['sapma'], r['onerilen_carpan']), ('kurulum', 100.0, 2.0))

    def test_musteri_veri_ozeti(self):
        import pandas as pd
        from teklif_veri import musteri_veri_ozeti
        pd.DataFrame({'a': range(10)}).to_excel(os.path.join(self.d, 'cari_dia_hazir.xlsx'), index=False)
        pd.DataFrame({'a': range(5)}).to_excel(os.path.join(self.d, 'stok_listesi.xlsx'), index=False)
        with pd.ExcelWriter(os.path.join(self.d, 'cari_OZET.xlsx')) as w:
            pd.DataFrame({'x': [1]}).to_excel(w, sheet_name='OZET')
            pd.DataFrame({'satir': [2, 2, 3]}).to_excel(w, sheet_name='DETAY')
        o = musteri_veri_ozeti(self.d)
        self.assertEqual((o['cari'], o['stok'], o['hata_orani']), (10, 5, round(2 * 100 / 15, 1)))

    def test_karsilastirma_ve_yeni_bolumler(self):
        g1 = {'teklif_no': 'RT-1', 'etiket': 'Ekonomik', 'musteri': 'Ş', 'masaustu': 3, 'mobil': 0,
              'kalemler': [{'kod': 'MAS2'}]}
        g2 = dict(g1, teklif_no='RT-2-R1', etiket='Önerilen', kalemler=[{'kod': 'MAS1'}], bakim='Standart')
        s1, s2 = self.h(**g1), self.h(**g2)
        b = teklif_cikti.bolumler(g2, s2, self.ayar)
        self.assertEqual(b[0][2], 'Revizyon 1 · Önerilen Teklif')
        basliklar = [x[1] for x in b if x[0] == 'bolum']
        for parca in ('Uyarlama Hizmeti', 'Ödeme Seçenekleri', 'Bakım ve Destek', 'Yıllık Toplam Maliyet'):
            self.assertTrue(any(parca in x for x in basliklar), parca)
        yol = teklif_cikti.karsilastirma_pdf(os.path.join(self.d, 'k.pdf'), [(g1, s1), (g2, s2)], self.ayar)
        with open(yol, 'rb') as f:
            self.assertTrue(f.read(4) == b'%PDF')
        teklif_cikti.karsilastirma_excel(os.path.join(self.d, 'k.xlsx'), [(g1, s1), (g2, s2)], self.ayar)

    def test_eski_pesin_indirim_ayari_tasinir(self):
        yol = os.path.join(self.d, 'eski.db')
        import sqlite3 as sq
        db = sq.connect(yol)
        db.execute('CREATE TABLE ayar (ad TEXT PRIMARY KEY, deger TEXT)')
        db.execute("INSERT INTO ayar VALUES ('pesin_indirim', '7')")
        db.commit()
        db.close()
        v = TeklifVeri(yol)
        self.assertEqual(v.kurallar()['odeme'][0]['indirim'], 7)
        v.db.close()


class KodYapistirTest(unittest.TestCase):
    P = {'DIA-URE-STD': {'aktif': 1}, 'EKO.MOBIL': {'aktif': 1}, 'ŞUBE1': {'aktif': 1}, 'ESKI': {'aktif': 0}}

    def test_alt_alta_ve_ayraclar(self):
        b, y, p = kod_listesi_coz('DIA-URE-STD\r\neko.mobil; şube1, ESKI\n\nYOKKOD', self.P)
        self.assertEqual(b, [('DIA-URE-STD', None), ('EKO.MOBIL', None), ('ŞUBE1', None)])
        self.assertEqual((y, p), (['YOKKOD'], ['ESKI']))

    def test_ayrac_farki_ve_adet(self):
        b, y, _ = kod_listesi_coz('dia ure std 3\nEKO_MOBIL', self.P)
        self.assertEqual(b, [('DIA-URE-STD', 3), ('EKO.MOBIL', None)])
        self.assertEqual(y, [])

    def test_excel_satirlari(self):
        metin = 'DIA-URE-STD\t2\tÜretim\t15000\nEKO.MOBIL\tMobil\t900\nBaşlık\tAd\n'
        b, y, _ = kod_listesi_coz(metin, self.P)
        self.assertEqual(b, [('DIA-URE-STD', 2), ('EKO.MOBIL', None)])
        self.assertEqual(y, ['Başlık'])

    def test_bos(self):
        self.assertEqual(kod_listesi_coz('', self.P), ([], [], []))


class DiaKampanyaTest(unittest.TestCase):
    """Kullanicinin DIA uye ekrani: 5 paket, E-Fatura'da 'Ekim 2026 E-Fatura' kampanyasi 7.800 TL.
    DIA: Uyelik giris 297.900, Ek +1 masaustu 17.100, Yillik bulut hizmet bedeli 99.300.
    Kampanya paketin kendisinde (DIA listesinden), teklife eklenince satirin Ind. % / Ind. Tutar'ina yazilir."""
    P = {'P12_PRKN_PRO': ('Masaüstü', 67600, 15, 15), 'P71_GENMUH': ('Masaüstü', 46400, 15, 15),
         'M031_POS_01': ('Diğer', 20700, 0, 50), 'M031_POS_1_14': ('Diğer', 132000, 0, 50),
         'M550_EFATURA_EARSIV_01': ('Diğer', 39000, 0, 15)}

    def setUp(self):
        self.pk = {k: {'kod': k, 'tur': t, 'ad': k, 'std_fiyat': f, 'ent_fiyat': 0, 'kull_yuzde': ku, 'ybh_yuzde': y,
                       'std_saat': 1, 'ent_saat': 1, 'aktif': 1} for k, (t, f, ku, y) in self.P.items()}
        self.pk['M550_EFATURA_EARSIV_01'].update(kamp_tutar=7800, kamp_aciklama='Ekim 2026 E-Fatura')
        self.ayar = dict(AYAR_VARSAYILAN, saatlik_ucret=100)

    def kalem(self, kod, **k):
        from teklif_hesap import paket_indirimi
        kalem = dict({'kod': kod}, **(paket_indirimi(self.pk[kod]) or {}))
        kalem.update(k)
        return kalem

    def h(self, **k):
        g = {'masaustu': 1, 'mobil': 0, 'kalemler': [self.kalem(c, **k.get(c, {})) for c in self.P]}
        return hesapla(g, self.pk, self.ayar)

    def test_paket_indirimi_satira_yazilir(self):
        from teklif_hesap import paket_indirimi
        self.assertEqual(paket_indirimi(self.pk['M550_EFATURA_EARSIV_01']),
                         {'ind_yuzde': 0, 'ind_tutar': 7800, 'kamp_aciklama': 'Ekim 2026 E-Fatura'})
        self.assertIsNone(paket_indirimi(self.pk['P12_PRKN_PRO']))

    def test_dia_ekrani_ile_ayni(self):
        s = self.h()
        self.assertEqual(s['dia'], {'uyelik_giris': 297900, 'kampanya_indirimi': 7800, 'ek_masaustu': 17100,
                                    'ek_mobil': 0, 'yillik_bulut': 99300})
        self.assertEqual(s['lisans']['net'], 297900)
        efat = s['satirlar'][-1]
        self.assertEqual((efat['kamp_ad'], efat['kamp_indirim'], efat['net']), ('Ekim 2026 E-Fatura', 7800, 31200))

    def test_indirim_satirda_degistirilebilir(self):
        s = self.h(M550_EFATURA_EARSIV_01={'ind_tutar': 0, 'ind_yuzde': 10})
        x = s['satirlar'][-1]
        self.assertEqual((x['indirim'], x['net'], x['kamp_indirim']), (3900, 35100, 3900))
        s = self.h(M550_EFATURA_EARSIV_01={'kamp_aciklama': ''})       # aciklama silinirse bizim indirimimiz sayilir
        self.assertEqual((s['satirlar'][-1]['kamp_indirim'], s['satirlar'][-1]['elle_indirim']), (0, 7800))

    def test_dia_kampanyasi_bizim_indirim_sayilmaz(self):
        self.ayar['max_toplam_indirim'] = 1
        s = self.h()
        self.assertFalse([k for k in s['kontrol']['liste'] if 'Toplam indirim' in k['mesaj']])

    def test_sifir_adam_saat_uyarisi(self):
        self.pk['P12_PRKN_PRO']['std_saat'] = 0
        s = hesapla({'masaustu': 1, 'mobil': 0, 'kalemler': [{'kod': 'P12_PRKN_PRO'}]}, self.pk, self.ayar,
                    kurallar={'hizmetler': [{'kod': 'kurulum', 'ad': 'Kurulum', 'hesap': 'PAKET_SAAT'}]})
        self.assertTrue([k for k in s['kontrol']['liste'] if 'Adam/saati 0' in k['mesaj']])

    def test_sayi_coz(self):
        for metin, beklenen in [('7.800', 7800), ('7 800,00', 7800), ('1.250.000', 1250000), ('20,5', 20.5),
                                ('20.5', 20.5), ('%15', 15), ('₺7.800', 7800), ('1,250.50', 1250.5)]:
            self.assertEqual(sayi_coz(metin), beklenen, metin)
        self.assertIsNone(sayi_coz('abc'))


class DiaListeIceAktarmaTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def liste(self, basliklar, satirlar, sayfa='Sheet'):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = sayfa
        ws.append(basliklar)
        for r in satirlar:
            ws.append(r)
        yol = os.path.join(self.d, 'dia.xlsx')
        wb.save(yol)
        return yol

    def test_aylik_listede_indirim_kolonlari_pakete_islenir(self):
        yol = self.liste(['Tür', 'Kodu', 'Adı', 'Std. Fiyat', 'Ent. Fiyat', 'İndirim %', 'İndirim Tutarı', 'Kampanya'],
                         [['Masaüstü', 'MAS1', 'Muhasebe', 1100, 2000, 10, 0, 'Ekim Muhasebe'],
                          ['Diğer', 'DIG1', 'E-Fatura', 800, 900, 0, 160, 'Ekim 2026 E-Fatura']])
        sonuc = self.v.paket_listesi_guncelle(yol)
        self.assertEqual((sonuc['kampanyali'], sonuc['indirim_kolonu']), (['DIG1', 'MAS1'], True))
        p = self.v.paketler()
        self.assertEqual((p['DIG1']['kamp_tutar'], p['DIG1']['kamp_aciklama']), (160, 'Ekim 2026 E-Fatura'))
        self.assertEqual((p['MAS1']['kamp_yuzde'], p['MAS1']['std_saat']), (10, 2))     # adam/saat korunur
        # Ertesi ay: kampanya bitti (0) -> indirim kalkar
        yol = self.liste(['Tür', 'Kodu', 'Adı', 'Std. Fiyat', 'Ent. Fiyat', 'İndirim %', 'İndirim Tutarı'],
                         [['Diğer', 'DIG1', 'E-Fatura', 800, 900, 0, 0]])
        self.v.paket_listesi_guncelle(yol)
        self.assertEqual((self.v.paketler()['DIG1']['kamp_tutar'], self.v.paketler()['DIG1']['kamp_aciklama']), (0, ''))

    def test_indirim_kolonu_olmayan_liste_elle_indirimi_korur(self):
        with self.v.db:
            self.v.db.execute("UPDATE paket SET kamp_tutar=50, kamp_aciklama='Elle' WHERE kod='DIG1'")
        # DIA'nin gercek listesi: basliklarda bosluklar, sayfa adi 'PaketListesi ' (sonda bosluk)
        yol = self.liste(['Türü    ', 'Kodu             ', 'Adı      ', ' Std. Fiyat', ' Ent. Fiyat', ' Kullanıcı %',
                          '     YBH %', 'Açıklama    ', 'Ürünler    ', 'Zorunlu     '],
                         [['Diğer', 'DIG1', 'E-Fatura', 850, 900, 0, 20, 'x', 'y', '[]']], sayfa='PaketListesi ')
        sonuc = self.v.paket_listesi_guncelle(yol)
        p = self.v.paketler()['DIG1']
        self.assertEqual((sonuc['indirim_kolonu'], p['std_fiyat'], p['kamp_tutar'], p['kamp_aciklama'], p['tur']),
                         (False, 850, 50, 'Elle', 'Diğer'))
        self.v.excelden_ice_aktar(yol)                  # motor ice aktarimi da sayfa adini bulur, indirimi korur
        self.assertEqual(self.v.paketler()['DIG1']['kamp_tutar'], 50)

    def test_indirim_basliklari_esnek(self):
        from teklif_veri import kamp_kolonu
        for baslik, beklenen in [('Kamp.İnd.%', 'kamp_yuzde'), ('İNDİRİM %', 'kamp_yuzde'), ('İskonto Oranı', 'kamp_yuzde'),
                                 ('Kamp.İnd.Tutar', 'kamp_tutar'), ('İndirim', 'kamp_tutar'),
                                 ('Kampanya İndirimi', 'kamp_tutar'), ('İndirim (TL)', 'kamp_tutar'),
                                 ('Kamp.Açıklama', 'kamp_aciklama'), ('Kampanya', 'kamp_aciklama'),
                                 (' Kullanıcı %', None), ('     YBH %', None), ('Açıklama', None), (' Std. Fiyat', None)]:
            self.assertEqual(kamp_kolonu(baslik), beklenen, baslik)

    def test_eski_kampanya_tablosu_pakete_tasinir(self):
        with self.v.db:
            self.v.db.execute("INSERT INTO paket_kampanya (ad, kodlar, yuzde, tutar, aktif) VALUES "
                              "('Ekim - DIG1', 'DIG*', 0, 160, 1)")
        self.v.db.close()
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        p = self.v.paketler()['DIG1']
        self.assertEqual((p['kamp_tutar'], p['kamp_aciklama']), (160, 'Ekim'))
        self.assertFalse(self.v.db.execute('SELECT 1 FROM paket_kampanya').fetchone())

    def test_adam_saati_olmayan_excel_eski_saatleri_silmez(self):
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(self.d, 'motor.xlsx'))
        del wb['ADAM SAAT']
        yol = os.path.join(self.d, 'saatsiz.xlsx')
        wb.save(yol)
        rapor = self.v.excelden_ice_aktar(yol)
        self.assertEqual(self.v.paketler()['MAS1']['std_saat'], 2)
        self.assertTrue(any('korundu' in r for r in rapor))


class TasarimTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))
        self.ayar = dict(self.v.ayarlar(), firma_telefon='0312 000', hazirlayan_ad='Ayşe Satış',
                         firma_banka='IBAN: TR00')
        self.g = {'teklif_no': 'RT-9', 'musteri': 'Örnek Ltd.', 'masaustu': 2, 'mobil': 0,
                  'musteri_bilgi': {'yetkili': 'Ali Veli'}, 'kalemler': [{'kod': 'MAS1'}, {'kod': 'DIG1'}]}
        self.s = hesapla(self.g, self.v.paketler(), self.ayar, kurallar=self.v.kurallar(gecmis=False))

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def test_hazir_tasarimlar_ve_varsayilan(self):
        self.assertEqual(sorted(self.v.tasarimlar()), ['Kurumsal', 'Kurumsal Kapaklı', 'Kurumsal Klasik',
                                                       'Kurumsal Şeritli', 'Modern', 'Sade'])
        self.assertEqual(self.v.varsayilan_tasarim_adi(), 'Kurumsal')
        t = self.v.tasarim('Modern')
        t['goster']['tco'] = False
        self.v.tasarim_kaydet('Benim', t, varsayilan=True)
        self.assertEqual(self.v.varsayilan_tasarim_adi(), 'Benim')
        self.assertFalse(self.v.tasarim()['goster']['tco'])
        self.assertEqual(self.v.tasarim('olmayan')['sablon'], 'Modern')     # bilinmeyen -> varsayilan
        self.v.tasarim_sil('Benim')
        self.assertIn(self.v.varsayilan_tasarim_adi(), self.v.tasarimlar())

    def test_alan_secimi_ve_metinler(self):
        import teklif_tasarim
        t = teklif_tasarim.hazir('Kurumsal')
        b = teklif_cikti.bolumler(self.g, self.s, self.ayar, t)
        turler = [x[0] for x in b]
        for tur in ('taraflar', 'paragraf', 'maddeler', 'imza', 'vurgu'):
            self.assertIn(tur, turler)
        self.assertIn('Sayın Ali Veli', next(x[1] for x in b if x[0] == 'paragraf'))
        lisans = next(x for x in b if x[0] == 'tablo')
        self.assertEqual(lisans[1], ['#', 'Paket / Modül', 'Kapsam', 'Liste Tutarı', 'İndirim', 'Net Tutar'])
        t['goster'].update(liste_indirim=False, kullanici_kolonu=False, on_yazi=False, taraflar=False, banka=False)
        b = teklif_cikti.bolumler(self.g, self.s, self.ayar, t)
        self.assertEqual(next(x for x in b if x[0] == 'tablo')[1], ['#', 'Paket / Modül', 'Net Tutar'])
        self.assertNotIn('taraflar', [x[0] for x in b])
        self.assertNotIn('Banka Bilgileri', [x[1] for x in b if x[0] == 'bolum'])

    def test_uc_sablon_pdf_ve_excel(self):
        import teklif_tasarim
        from openpyxl import load_workbook
        for ad in teklif_tasarim.SABLONLAR:
            t = teklif_tasarim.hazir(ad)
            pdf = teklif_cikti.pdf_olustur(os.path.join(self.d, ad + '.pdf'), self.g, self.s, self.ayar, t)
            with open(pdf, 'rb') as f:
                self.assertTrue(f.read(4) == b'%PDF', ad)
            xl = teklif_cikti.excel_olustur(os.path.join(self.d, ad + '.xlsx'), self.g, self.s, self.ayar, t)
            hucreler = [c.value for r in load_workbook(xl).active.iter_rows() for c in r if c.value is not None]
            self.assertIn('Örnek Ltd.', hucreler)
            self.assertIn(self.s['genel_toplam'], hucreler)

    def test_bozuk_tasarim_varsayilana_doner(self):
        import teklif_tasarim
        t = teklif_tasarim.tamamla({'sablon': 'Yok', 'ana_renk': 'mavi', 'yazi_boyutu': 'x', 'goster': {'tco': 0}})
        self.assertEqual((t['sablon'], t['ana_renk'], t['yazi_boyutu']), ('Kurumsal', '#1F3A5F', 9))
        self.assertFalse(t['goster']['tco'])
        self.assertTrue(t['goster']['imza'])


class KesifVeAsistanTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))
        self.k = self.v.kurallar(gecmis=False)
        self.ayar = dict(self.v.ayarlar(), max_toplam_indirim=20)

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def h(self, **g):
        return hesapla(dict({'musteri': 'A', 'masaustu': 2, 'mobil': 0, 'kalemler': [{'kod': 'MAS1'}]}, **g),
                       self.v.paketler(), self.ayar, kurallar=self.k)

    def test_standart_cevaplar_ve_gruplar(self):
        from teklif_hesap import standart_cevaplar
        c = standart_cevaplar(self.k['sorular'])
        self.assertEqual((c['sube_sayisi'], c['uretim'], c['e_donusum']), (1, False, True))
        self.assertNotIn('cari_sayisi', c)
        self.assertTrue(all(s.get('grup') for s in self.k['sorular']))

    def test_soru_duzenleme_ve_varsayilan(self):
        sorular = self.v.tablo_oku('kesif')
        for s in sorular:
            if s['kod'] == 'sube_sayisi':
                s.update(soru='Mağaza sayısı', varsayilan='3')
        self.v.tablo_yaz('kesif', sorular)
        from teklif_hesap import standart_cevaplar
        k = self.v.kurallar(gecmis=False)
        self.assertEqual([s['soru'] for s in k['sorular'] if s['kod'] == 'sube_sayisi'], ['Mağaza sayısı'])
        self.assertEqual(standart_cevaplar(k['sorular'])['sube_sayisi'], 3)

    def test_profiller(self):
        self.assertIn('Üretim firması', self.k['profiller'])
        self.v.kesif_profil_kaydet('Benim', {'sube_sayisi': 4, 'uretim': True, 'bos': ''})
        self.assertEqual(self.v.kesif_profilleri()['Benim'], {'sube_sayisi': 4, 'uretim': True})
        self.v.kesif_profil_sil('Benim')
        self.assertNotIn('Benim', self.v.kesif_profilleri())

    def test_etki_metni(self):
        from teklif_hesap import soru_etki_metni
        s = {s['kod']: s for s in self.k['sorular']}
        self.assertEqual(soru_etki_metni(s['uretim'], {'kurulum': 'Kurulum'}), 'Evet ise Kurulum: +24 saat')
        self.assertIn('× 1.15', soru_etki_metni(s['it_personeli_yok']))
        self.assertTrue(soru_etki_metni(s['yeni_musteri']).startswith('Sadece bilgi'))

    def test_asistan_ihtiyac_pazarlik_hazirlik(self):
        s = self.h(cevaplar={'uretim': True, 'e_donusum': True})
        a = s['asistan']
        basliklar = [o['baslik'] for o in a['oneriler']]
        self.assertIn('Üretim paketi yok', basliklar)
        self.assertIn('e-Dönüşüm paketi yok', basliklar)
        s2 = self.h(cevaplar={'uretim': True}, kalemler=[{'kod': 'MAS1'}, {'kod': 'URE'}])
        self.assertNotIn('Üretim paketi yok', [o['baslik'] for o in s2['asistan']['oneriler']])
        p = a['pazarlik']
        self.assertAlmostEqual(p['taban'], p['liste'] * 0.8, places=1)
        self.assertTrue(p['taban'] <= p['hedef'] <= p['teklif'])
        self.assertEqual(len(a['takip']), 3)
        self.assertIn(('Karar verici (yetkili)', False), a['hazirlik'])
        from teklif_hesap import satis_brifingi
        metin = satis_brifingi({'musteri': 'A'}, s)
        self.assertIn('SATIŞ BRİFİNGİ', metin)
        self.assertIn('Taviz sırası', metin)


class GorunumDovizDilTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))
        self.ayar = self.v.ayarlar()
        self.g = {'teklif_no': 'RT-1', 'musteri': 'Örnek', 'masaustu': 2, 'mobil': 0, 'bakim': 'Standart',
                  'kalemler': [{'kod': 'MAS1'}, {'kod': 'DIG1'}]}

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def h(self, **g):
        return hesapla(dict(self.g, **g), self.v.paketler(), self.ayar, kurallar=self.v.kurallar(gecmis=False))

    def test_uyarlama_ve_ozet_gorunumu(self):
        import teklif_tasarim
        s = self.h()
        t = teklif_tasarim.hazir('Kurumsal')
        t.update(uyarlama_gorunum='Özet', ozet_gorunum='Kısa')
        b = teklif_cikti.bolumler(self.g, s, self.ayar, t)
        self.assertEqual(sum(1 for x in b if x[0] == 'tablo' and 'Hizmet' in x[1]), 0)     # uyarlama tablosu yok
        vurgu = next(x for x in b if x[0] == 'vurgu')
        self.assertIn('Kurulum ve parametre ayarları', vurgu[3])
        ozet = next(x for x in b if x[0] == 'ozet')
        self.assertEqual([r[0] for r in ozet[1]], ['NET TOPLAM (KDV Hariç)', 'KDV (%20)', 'GENEL TOPLAM (KDV Dahil)'])
        t['ozet_gorunum'] = 'Gizle'
        self.assertFalse([x for x in teklif_cikti.bolumler(self.g, s, self.ayar, t) if x[0] == 'ozet'])
        t['uyarlama_gorunum'] = 'Kalem listesi'
        tablo = next(x for x in teklif_cikti.bolumler(self.g, s, self.ayar, t) if x[0] == 'tablo' and
                     x[1][0] == 'Hizmet')
        self.assertNotIn('Dayanak', ' '.join(r[1] for r in tablo[2]))

    def test_bakim_aylik_gosterim(self):
        import teklif_tasarim
        s = self.h()
        t = teklif_tasarim.hazir('Kurumsal')
        bakim = next(x for x in teklif_cikti.bolumler(self.g, s, self.ayar, t) if x[0] == 'tablo' and
                     'Aylık (KDV Hr.)' in x[1])
        self.assertNotIn('Yıllık (KDV Hr.)', bakim[1])
        t['goster']['bakim'] = False
        self.assertFalse([x for x in teklif_cikti.bolumler(self.g, s, self.ayar, t)
                          if x[0] == 'bolum' and 'Bakım' in x[1]])

    def test_doviz_ve_ingilizce(self):
        import teklif_tasarim
        g = dict(self.g, doviz='USD', kur=40, dil='EN')
        s = self.h(**g)
        b = teklif_cikti.bolumler(g, s, self.ayar, teklif_tasarim.hazir('Kurumsal'))
        self.assertEqual(b[1][1], 'DIA ERP LICENSE AND SERVICES PROPOSAL')
        ozet = next(x for x in b if x[0] == 'ozet')
        self.assertEqual((ozet[3], ozet[4]), ('USD', 'EN'))
        self.assertAlmostEqual(ozet[1][-1][1], round(s['genel_toplam'] / 40, 2))
        self.assertIn('Licenses', [x[1][3:] for x in b if x[0] == 'bolum'])
        self.assertTrue(any('1 USD = 40' in m for x in b if x[0] == 'maddeler' for m in x[1]))
        pdf = teklif_cikti.pdf_olustur(os.path.join(self.d, 'en.pdf'), g, s, self.ayar)
        self.assertTrue(os.path.getsize(pdf) > 1000)
        self.assertEqual(teklif_cikti.bicim(1234.5, 'USD', 'EN'), '1,234.50 USD')
        self.assertEqual(teklif_cikti.doviz_bilgisi({'doviz': 'EUR'}), ('TL', 1.0))       # kur yoksa TL

    def test_hizmet_adi_teklife_ozel(self):
        s = self.h(hizmet_metin={'kurulum': {'ad': 'Sistem kurulumu', 'kapsam': 'Kısa kapsam'}})
        k = next(x for x in s['hizmet']['kalemler'] if x['kod'] == 'kurulum')
        self.assertEqual((k['ad'], k['kapsam'], k['ozel_ad']), ('Sistem kurulumu', 'Kısa kapsam', True))
        b = teklif_cikti.bolumler(self.g, s, self.ayar)
        tablo = next(x for x in b if x[0] == 'tablo' and x[1][0] == 'Hizmet')
        self.assertEqual(tablo[2][0][0], 'Sistem kurulumu')

    def test_bos_teklifte_uyarlama_yok(self):
        s = self.h(kalemler=[], cevaplar={'e_donusum': True, 'sube_sayisi': 4})
        self.assertEqual(s['uyarlama']['saat'], 0)

    def test_tum_sablonlar(self):
        import teklif_tasarim
        s = self.h()
        for ad in teklif_tasarim.SABLONLAR:
            t = teklif_tasarim.tamamla({'sablon': ad})
            pdf = teklif_cikti.pdf_olustur(os.path.join(self.d, 'x.pdf'), self.g, s, self.ayar, t)
            self.assertTrue(os.path.getsize(pdf) > 1000, ad)


class SatisSureciTest(unittest.TestCase):
    """Asama, takip, aktivite, gunluk isler, gecmisten ogrenme ve rapor."""
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = self.tmp.name
        motor_excel(os.path.join(self.d, 'motor.xlsx'))
        self.v = TeklifVeri(os.path.join(self.d, 'teklif.db'))
        self.v.excelden_ice_aktar(os.path.join(self.d, 'motor.xlsx'))
        self.v.ayar_yaz('hazirlayan_ad', 'Ayşe')

    def tearDown(self):
        self.v.db.close()
        self.tmp.cleanup()

    def kaydet(self, musteri, tarih, ind=0, bakim=''):
        g = {'musteri': musteri, 'masaustu': 2, 'mobil': 0, 'tarih': tarih, 'kampanya': 'X', 'kampanya_yuzde': ind,
             'kalemler': [{'kod': 'MAS1'}], 'bakim': bakim}
        s = hesapla(g, self.v.paketler(), self.v.ayarlar(), kurallar=self.v.kurallar())
        return self.v.teklif_kaydet(g, s), g, s

    def test_kayitta_asama_takip_hazirlayan(self):
        from datetime import datetime, timedelta
        no, g, s = self.kaydet('A', '2026-01-10')
        t = self.v.teklif_sureci(no)
        self.assertEqual((t['asama'], t['takip_tarihi'], t['hazirlayan']), ('Teklif verildi', '2026-01-12', 'Ayşe'))
        self.v.asama_yaz(no, 'Müzakere')
        self.v.teklif_kaydet(dict(g, teklif_no=no), s, no)            # tekrar kayit sureci bozmaz
        self.assertEqual(self.v.teklif_sureci(no)['asama'], 'Müzakere')
        self.v.teklif_sonuc_yaz(no, 'KAZANDI')
        t = self.v.teklif_sureci(no)
        self.assertEqual(t['asama'], 'Kazandı')
        self.assertEqual(t['takip_tarihi'], (datetime.now() + timedelta(days=30)).strftime('%Y-%m-%d'))
        self.assertEqual([a['tur'] for a in self.v.aktiviteler(no)], ['Sonuç', 'Aşama'])

    def test_aktivite_ve_gunluk_isler(self):
        no, _, _ = self.kaydet('Gecikmiş', '2026-01-01')
        no2, _, _ = self.kaydet('İleri', '2026-01-01')
        self.v.aktivite_ekle(no2, 'Arama', 'demo istendi', '2099-01-01')
        isler = self.v.gunluk_isler('2026-01-05')
        self.assertEqual([(x['no'], x['oncelik']) for x in isler if x['no'] == no], [(no, 0)])
        self.assertIn('gecikti', isler[0]['is'])
        self.assertNotIn(no2, [x['no'] for x in isler if x['is'].startswith('Takip')])
        self.assertEqual(self.v.aktiviteler(no2)[0]['metin'], 'demo istendi')

    def test_senaryolar_ve_sonraki_adim(self):
        self.assertTrue([x for x in self.v.senaryolar() if x['tur'] == 'ASAMA' and x['baslik'] == 'Müzakere'])
        g = {'musteri': 'B', 'masaustu': 2, 'mobil': 0, 'kalemler': [{'kod': 'MAS1'}], 'asama': 'Müzakere',
             'teklif_no': 'X'}
        a = hesapla(g, self.v.paketler(), dict(self.v.ayarlar(), max_toplam_indirim=20),
                    kurallar=self.v.kurallar())['asistan']
        self.assertEqual(a['asama'], 'Müzakere')
        self.assertTrue(any('taban' in x for x in a['sonraki_adim']))
        self.assertTrue(any('Rakip' in i[0] or 'Fiyat' in i[0] for i in a['itirazlar']))
        sen = self.v.tablo_oku('senaryo')
        for x in sen:
            if x['baslik'] == 'Müzakere':
                x['metin'] = 'Özel adım {hedef}'
        self.v.tablo_yaz('senaryo', sen)
        a = hesapla(g, self.v.paketler(), self.v.ayarlar(), kurallar=self.v.kurallar())['asistan']
        self.assertTrue(any(x.startswith('Özel adım ') and '{' not in x for x in a['sonraki_adim']))

    def test_gecmisten_ogrenme(self):
        for i, (ind, sonuc, bakim) in enumerate([(0, 'KAYBETTİ', ''), (2, 'KAYBETTİ', ''), (3, 'KAYBETTİ', ''),
                                                (6, 'KAZANDI', 'Standart'), (7, 'KAZANDI', 'Standart'),
                                                (8, 'KAZANDI', 'Standart'), (9, 'KAYBETTİ', '')]):
            no, _, _ = self.kaydet('M%d' % i, '2026-01-%02d' % (i + 1), ind, bakim)
            self.v.teklif_sonuc_yaz(no, sonuc, 'fiyat yüksek' if sonuc == 'KAYBETTİ' else '')
        g = {'musteri': 'M1', 'masaustu': 2, 'mobil': 0, 'kalemler': [{'kod': 'MAS1'}]}
        og = hesapla(g, self.v.paketler(), self.v.ayarlar(), kurallar=self.v.kurallar())['asistan']['ogrenilen']
        metin = ' '.join(og)
        self.assertIn('daha önce 1 teklif', metin)
        self.assertIn('%5-%10', metin)
        self.assertIn('Bakım anlaşması eklenen tekliflerde kazanma %100', metin)
        self.assertIn('"fiyat"', metin)

    def test_rapor(self):
        import teklif_rapor
        a, _, _ = self.kaydet('A', '2026-02-01')
        b, _, _ = self.kaydet('B', '2026-02-10')
        self.kaydet('C', '2026-03-01')
        self.v.teklif_sonuc_yaz(a, 'KAZANDI')
        self.v.teklif_sonuc_yaz(b, 'KAYBETTİ', 'rakip')
        r = teklif_rapor.rapor(teklif_rapor.filtrele(self.v.rapor_verisi(), '2026-01-01'),   # ESKI-1 haric
                               self.v.aktiviteler())
        self.assertEqual((r['kpi']['teklif'], r['kpi']['kazanilan'], r['kpi']['kaybedilen'], r['kpi']['kazanma_orani']),
                         (3, 1, 1, 50))
        self.assertEqual([x['ay'] for x in r['aylik']], ['2026-02', '2026-03'])
        self.assertEqual(r['satisci'][0]['hazirlayan'], 'Ayşe')
        self.assertEqual(r['nedenler'][0]['neden'], 'rakip')
        self.assertEqual(r['paketler'][0], {'kod': 'MAS1', 'teklif': 3, 'kazanilan': 1})
        self.assertEqual(len(teklif_rapor.filtrele(self.v.rapor_verisi(), '2026-02-05')), 2)
        teklif_rapor.excel_yaz(os.path.join(self.d, 'r.xlsx'), r)


if __name__ == '__main__':
    unittest.main()
